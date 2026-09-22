from openpyxl import load_workbook
from pathlib import Path

path = Path(r'C:\Users\施昊lty\Desktop\文案\海洋馆募捐调查委托_杏仁网加密文件已填入.xlsx')
wb = load_workbook(path, read_only=True, data_only=False)
for ws in wb.worksheets:
    print(f'\n=== {ws.title} {ws.max_row}x{ws.max_column} ===')
    if ws.title == '杏仁网加密文件总表':
        for idx, row in enumerate(ws, 1):
            vals = [f'{cell.column_letter}:{str(cell.value)[:700]}' for cell in row if cell.value is not None]
            if vals:
                print(f'{idx}: ' + ' | '.join(vals))
    else:
        for idx, row in enumerate(ws, 1):
            vals = [str(cell.value) for cell in row if cell.value is not None]
            if any(x in ' '.join(vals) for x in ('贾老板','钱蔚东','贾剑明','基金会搜索引擎','蓝鲸贸易')):
                print(f'{idx}: ' + ' | '.join(x[:400] for x in vals))
