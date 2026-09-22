from pathlib import Path
from openpyxl import load_workbook
from html.parser import HTMLParser

class TagCheck(HTMLParser):
    def __init__(self):
        super().__init__()
        self.tags = []
    def handle_starttag(self, tag, attrs):
        self.tags.append(tag)

root = Path(r'F:\first!!!')
src = Path(r'C:\Users\施昊lty\Desktop\文案\海洋馆募捐调查委托_杏仁网加密文件已填入.xlsx')
out = root / 'outputs' / '钱蔚东线' / '海洋馆募捐调查委托_钱蔚东证据重写.xlsx'
a = load_workbook(src, data_only=False)
b = load_workbook(out, data_only=False)
assert a.sheetnames == b.sheetnames, (a.sheetnames, b.sheetnames)
allowed = {('杏仁网加密文件总表', f'{c}4') for c in 'BCDEF'} | {('杏仁网社交账号', 'E43')}
changed = []
for name in a.sheetnames:
    sa, sb = a[name], b[name]
    for row in sa:
        for ca in row:
            cb = sb[ca.coordinate]
            if ca.value != cb.value:
                changed.append((name, ca.coordinate))
                assert (name, ca.coordinate) in allowed, (name, ca.coordinate)
    assert list(sa.merged_cells.ranges) == list(sb.merged_cells.ranges), name
print('Changed cells:', changed)
assert set(changed) == allowed
target = b['杏仁网加密文件总表']
text = target['F4'].value
assert '2025.1.5' not in text and '2025.1.18' not in text
assert '蓝鲸贸易' in text and target['D4'].value == '20250103'
assert '¥1,860,000' in text and '¥1,768,000' in text
assert 780000 + 460000 + 620000 == 1860000
assert 742000 + 437000 + 589000 == 1768000
for file in ('bluewhale-evidence.html','mid/middle-bluewhale.html','socialaccount2.html','socialaccount3.html'):
    source = (root/file).read_text(encoding='utf-8')
    parser = TagCheck()
    parser.feed(source)
    assert 'html' in parser.tags and 'title' in parser.tags, file
    print(file, 'HTML tags:', len(parser.tags))
assert (root/'mid/middle-bluewhale.html').read_text(encoding='utf-8').count('bluewhale-evidence.html') == 1
assert 'keyword === "蓝鲸贸易"' in (root/'maiyingfoundation.html').read_text(encoding='utf-8')
assert 'archive: "20250104"' in (root/'socialaccount2.html').read_text(encoding='utf-8')
assert 'investigation: "20250103"' in (root/'socialaccount3.html').read_text(encoding='utf-8')
print('Verification passed')
